import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

doc = docx.Document()

for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(14)

md_path = r"t:\CLAUDE\Hop_Den_Tri_Thuc\outputs\2026-07-24_Can_Bo_So_Sai\YouTube\06_kich_ban_hoan_thien.md"
docx_path = r"t:\CLAUDE\Hop_Den_Tri_Thuc\outputs\2026-07-24_Can_Bo_So_Sai\YouTube\06_kich_ban_hoan_thien.docx"

with open(md_path, "r", encoding="utf-8") as f:
    lines = f.readlines()

for line in lines:
    line = line.rstrip('\n')
    if not line.strip():
        continue
    
    if line.strip() == '---':
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run('────────────────────────────────────────')
        run.font.color.rgb = RGBColor(200, 200, 200)
        continue
        
    if line.startswith('# '):
        p = doc.add_paragraph()
        r = p.add_run(line[2:].strip())
        r.bold = True
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor(192, 0, 0)
        continue

    if line.startswith('## '):
        p = doc.add_paragraph()
        r = p.add_run(line[3:].strip())
        r.bold = True
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor(0, 102, 204)
        continue

    if line.startswith('### '):
        p = doc.add_paragraph()
        r = p.add_run(line[4:].strip())
        r.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(204, 102, 0)
        continue

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    
    # Process basic bold tags
    parts = line.split("**")
    is_bold = False
    for part in parts:
        if part:
            clean_part = part.replace('<red>', '').replace('</red>', '').replace('<green>', '').replace('</green>', '').replace('<blue>', '').replace('</blue>', '').replace('<orange>', '').replace('</orange>', '').replace('<underline>', '').replace('</underline>', '')
            r = p.add_run(clean_part)
            if is_bold:
                r.bold = True
                r.font.color.rgb = RGBColor(192, 0, 0)
        is_bold = not is_bold

doc.save(docx_path)
print("Saved complete docx with 10 shorts at bottom:", docx_path)

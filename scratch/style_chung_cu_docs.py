from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"T:\CLAUDE\Hop_Den_Tri_Thuc")
OUT_DIR = ROOT / "drafts" / "2026-08-21_chung_cu_het_han"

INK = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
CALLOUT = "F4F6F9"


def set_typeface(font, name="Calibri", size=None, color=None, bold=None, italic=None):
    font.name = name
    font._element.rPr.rFonts.set(qn("w:ascii"), name)
    font._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    font._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        font.size = Pt(size)
    if color is not None:
        font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic


def set_style(style, *, size, color, bold=False, italic=False, before=0, after=0, line=1.1):
    set_typeface(style.font, size=size, color=color, bold=bold, italic=italic)
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def get_or_add_style(doc, name):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    for child in list(p_pr):
        if child.tag == qn("w:shd"):
            p_pr.remove(child)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def apply_run_format(paragraph, *, size, color, bold=None, italic=None):
    for run in paragraph.runs:
        set_typeface(run.font, size=size, color=color)
        if bold is not None:
            run.bold = bold
        if italic is not None:
            run.italic = italic


def replace_text(paragraph, value):
    p_pr = paragraph._p.pPr
    paragraph._p.clear_content()
    if p_pr is not None and paragraph._p.pPr is None:
        paragraph._p.insert(0, p_pr)
    paragraph.add_run(value)


def setup_document(doc, title):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    set_style(normal, size=11, color="000000", after=6, line=1.1)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    title_style = doc.styles["Title"]
    set_style(title_style, size=20, color=INK, bold=True, before=0, after=12, line=1.1)

    h1 = doc.styles["Heading 1"]
    set_style(h1, size=16, color=BLUE, bold=True, before=16, after=8, line=1.1)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    set_style(h2, size=13, color=BLUE, bold=True, before=12, after=6, line=1.1)
    h2.paragraph_format.keep_with_next = True

    h3 = doc.styles["Heading 3"]
    set_style(h3, size=12, color=DARK_BLUE, bold=True, before=8, after=4, line=1.1)
    h3.paragraph_format.keep_with_next = True

    quote = get_or_add_style(doc, "Hook Quote")
    quote.base_style = normal
    set_style(quote, size=12, color=INK, bold=True, italic=True, before=2, after=6, line=1.1)
    quote.paragraph_format.left_indent = Inches(0.22)
    quote.paragraph_format.right_indent = Inches(0.12)

    bridge = get_or_add_style(doc, "Cau cau noi")
    bridge.base_style = normal
    set_style(bridge, size=11, color=INK, bold=False, italic=False, before=2, after=9, line=1.1)
    bridge.paragraph_format.left_indent = Inches(0.12)
    bridge.paragraph_format.right_indent = Inches(0.12)

    list_number = doc.styles["List Number"]
    set_style(list_number, size=11, color="000000", after=8, line=1.167)
    list_number.paragraph_format.left_indent = Inches(0.5)
    list_number.paragraph_format.first_line_indent = Inches(-0.25)

    doc.core_properties.title = title
    doc.core_properties.subject = "Bài viết và gói SEO"
    doc.core_properties.author = "Hộp Đen Tri Thức"


def style_article(path):
    doc = Document(path)
    setup_document(doc, "Chung cư hết hạn có mất nhà không?")

    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        if index == 0:
            paragraph.style = doc.styles["Title"]
            apply_run_format(paragraph, size=20, color=INK, bold=True)
        elif text in {"Gợi ý Hook mở đầu", "Bài viết khoảng 200 từ - dùng Hook 1"}:
            paragraph.style = doc.styles["Heading 1"]
            apply_run_format(paragraph, size=16, color=BLUE, bold=True)
        elif text.startswith("Phương án "):
            paragraph.style = doc.styles["Heading 2"]
            apply_run_format(paragraph, size=13, color=BLUE, bold=True)
        elif text.startswith("Câu cầu nối nêu vấn đề:"):
            paragraph.style = doc.styles["Cau cau noi"]
            shade_paragraph(paragraph, CALLOUT)
            apply_run_format(paragraph, size=11, color=INK)
        elif text.startswith("Chung cư có ngày"):
            paragraph.style = doc.styles["Hook Quote"]
            apply_run_format(paragraph, size=12, color=INK, bold=True, italic=True)
        elif text.startswith("Cái cần làm rõ là:"):
            paragraph.style = doc.styles["Cau cau noi"]
            shade_paragraph(paragraph, CALLOUT)
            apply_run_format(paragraph, size=11, color=INK)
        else:
            paragraph.style = doc.styles["Normal"]
            apply_run_format(paragraph, size=11, color="000000")

    doc.save(path)


def style_seo(path):
    doc = Document(path)
    setup_document(doc, "Gói SEO: Chung cư hết hạn có mất nhà không?")

    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        if index == 0:
            paragraph.style = doc.styles["Title"]
            apply_run_format(paragraph, size=20, color=INK, bold=True)
        elif text in {"10 tiêu đề ngắn gọn, thu hút", "Mô tả chuẩn SEO", "Bộ thẻ Hashtag tối ưu tìm kiếm"}:
            paragraph.style = doc.styles["Heading 1"]
            apply_run_format(paragraph, size=16, color=BLUE, bold=True)
        elif text[0].isdigit() and ". " in text[:4]:
            item_text = text.split(". ", 1)[1]
            replace_text(paragraph, item_text)
            paragraph.style = doc.styles["List Number"]
            apply_run_format(paragraph, size=11, color="000000")
        else:
            paragraph.style = doc.styles["Normal"]
            apply_run_format(paragraph, size=11, color="000000")

    doc.save(path)


if __name__ == "__main__":
    style_article(OUT_DIR / "06_bai_viet_200_tu.docx")
    style_seo(OUT_DIR / "06b_goi_seo.docx")
    print("[OK] Applied standard_business_brief tokens and social-post formatting.")

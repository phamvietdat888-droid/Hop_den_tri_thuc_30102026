import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

# 1. Generate 06f_shorts_scripts.docx
doc1 = docx.Document()
for section in doc1.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

title_p = doc1.add_paragraph()
title_run = title_p.add_run("10 ĐOẠN KỊCH BẢN YOUTUBE SHORTS (MỖI ĐOẠN ~120 TỪ)")
title_run.bold = True
title_run.font.size = Pt(18)
title_run.font.color.rgb = RGBColor(192, 0, 0)
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

sub_p = doc1.add_paragraph()
sub_run = sub_p.add_run("Chủ đề: Cán Bộ Sợ Sai & Chiếc Chìa Khóa Nghị Định 73\nKênh: Hộp Đen Tri Thức | Tiêu Đề Viral High-CTR + Bối Cảnh Độc Lập Rõ Nét")
sub_run.font.size = Pt(11)
sub_run.font.italic = True
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc1.add_paragraph("-" * 50)

with open(r"t:\CLAUDE\Hop_Den_Tri_Thuc\outputs\2026-07-24_Can_Bo_So_Sai\YouTube_Shorts\06f_10_doan_shorts_sieu_thu_hut.md", "r", encoding="utf-8") as f:
    text = f.read()

sections_text = text.split("### ĐOẠN ")

for sec in sections_text[1:]:
    lines = sec.strip().split("\n")
    header_line = lines[0]
    
    h = doc1.add_heading(f"ĐOẠN {header_line}", level=2)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0, 102, 204)
        run.bold = True
    
    for line in lines[1:]:
        line = line.strip()
        if not line:
            continue
        if "Tiêu đề Shorts Viral:" in line:
            p = doc1.add_paragraph()
            r1 = p.add_run("🎬 TIÊU ĐỀ SHORTS VIRAL: ")
            r1.bold = True
            r1.font.color.rgb = RGBColor(204, 102, 0)
            clean_title = line.split("Tiêu đề Shorts Viral:")[1].strip()
            r2 = p.add_run(clean_title)
            r2.bold = True
            r2.font.size = Pt(15)
        elif line.startswith("* **Thời lượng dự kiến:**"):
            p = doc1.add_paragraph()
            r1 = p.add_run("⏱️ Thời lượng: ")
            r1.bold = True
            r1.font.color.rgb = RGBColor(46, 139, 87)
            p.add_run(line.replace("* **Thời lượng dự kiến:**", "").strip())
        elif line.startswith("**Nội dung kịch bản:**"):
            p = doc1.add_paragraph()
            r = p.add_run("🗣️ NỘI DUNG KỊCH BẢN (~120 TỪ):")
            r.bold = True
            r.font.color.rgb = RGBColor(128, 0, 128)
        elif line.startswith("* **💡 Chốt hạ / CTA:**") or line.startswith("* **Chốt hạ / CTA:**"):
            p = doc1.add_paragraph()
            r1 = p.add_run("💡 CHỐT HẠ / CTA: ")
            r1.bold = True
            r1.font.color.rgb = RGBColor(192, 0, 0)
            r2 = p.add_run(line.replace("* **💡 Chốt hạ / CTA:**", "").replace("* **Chốt hạ / CTA:**", "").strip())
            r2.bold = True
            r2.font.italic = True
        elif not line.startswith("---"):
            p = doc1.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(6)
            
            parts = line.split("**")
            is_bold = False
            for part in parts:
                if part:
                    run = p.add_run(part)
                    if is_bold:
                        run.bold = True
                        run.font.color.rgb = RGBColor(192, 0, 0)
                is_bold = not is_bold

    doc1.add_paragraph("")

out_dir = r"t:\CLAUDE\Hop_Den_Tri_Thuc\outputs\2026-07-24_Can_Bo_So_Sai\YouTube_Shorts"
os.makedirs(out_dir, exist_ok=True)
doc1.save(os.path.join(out_dir, "06f_shorts_scripts.docx"))
doc1.save(r"t:\CLAUDE\Hop_Den_Tri_Thuc\outputs\10_doan_shorts_so_sai.docx")

# 2. Generate 06_kich_ban_hoan_thien.docx
doc2 = docx.Document()
for section in doc2.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

style = doc2.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(14)

md_path = r"t:\CLAUDE\Hop_Den_Tri_Thuc\outputs\2026-07-24_Can_Bo_So_Sai\YouTube\06_kich_ban_hoan_thien.md"
docx_path = r"t:\CLAUDE\Hop_Den_Tri_Thuc\outputs\2026-07-24_Can_Bo_So_Sai\YouTube\06_kich_ban_hoan_thien.docx"

with open(md_path, "r", encoding="utf-8") as f:
    lines = f.readlines()

for line in lines:
    line = line.rstrip('\n')
    if not line.strip():
        continue
    
    if line.strip() == '---':
        p = doc2.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run('────────────────────────────────────────')
        run.font.color.rgb = RGBColor(200, 200, 200)
        continue
        
    if line.startswith('# '):
        p = doc2.add_paragraph()
        r = p.add_run(line[2:].strip())
        r.bold = True
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor(192, 0, 0)
        continue

    if line.startswith('## '):
        p = doc2.add_paragraph()
        r = p.add_run(line[3:].strip())
        r.bold = True
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor(0, 102, 204)
        continue

    if line.startswith('### '):
        p = doc2.add_paragraph()
        r = p.add_run(line[4:].strip())
        r.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(204, 102, 0)
        continue

    p = doc2.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    
    parts = line.split("**")
    is_bold = False
    for part in parts:
        if part:
            clean_part = part.replace('<red>', '').replace('</red>', '').replace('<green>', '').replace('</green>', '').replace('<blue>', '').replace('</blue>', '').replace('<orange>', '').replace('</orange>', '').replace('<underline>', '').replace('</underline>', '')
            r = p.add_run(clean_part)
            if is_bold:
                r.bold = True
                r.font.color.rgb = RGBColor(192, 0, 0)
        is_bold = not is_bold

doc2.save(docx_path)
print("Regenerated all docx files with viral titles.")

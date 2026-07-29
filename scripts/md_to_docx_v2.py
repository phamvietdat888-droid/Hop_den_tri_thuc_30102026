"""
md_to_docx_v2.py — Chuyển đổi kịch bản Markdown (Audio Clean) → Word (.docx)
Chuẩn định dạng: Tham chiếu từ 2026-04-22_Trump_5_Thang_Lich_Su/YouTube/06_kich_ban_hoan_thien.docx
Quy chuẩn: .shared/format_output.md (mục QUY CHUẨN FILE WORD)

Sử dụng: python md_to_docx_v2.py <input.md> <output.docx>
"""

import html
import os
import re
import sys

try:
    import docx
    from docx.shared import Pt, RGBColor, Cm, Emu
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False
    class RGBColor:
        def __init__(self, r, g, b):
            self.rgb = (r, g, b)

# === MÀU SẮC CHUẨN (từ file mẫu) ===
COLOR_MAP = {
    'red': RGBColor(204, 0, 0),       # CC0000 — Nhấn mạnh, bức xúc
    'green': RGBColor(0, 128, 0),      # 008000 — Dữ liệu, con số
    'blue': RGBColor(0, 102, 204),     # 0066CC — Mỉa mai, châm biếm
    'orange': RGBColor(230, 80, 0),    # E65000 — Câu hỏi, tò mò
}
COLOR_BLACK = RGBColor(0, 0, 0)
COLOR_DIVIDER = RGBColor(200, 200, 200)  # C8C8C8
DIVIDER_CHAR = '─' * 40


def parse_inline(text):
    """
    Parse inline formatting: <red>, <green>, <blue>, <orange>, <underline>,
    **bold**, *italic* — trả về list các segment:
    [{'text': ..., 'color': RGBColor|None, 'bold': bool, 'italic': bool, 'underline': bool}, ...]
    """
    segments = []
    
    pattern = re.compile(
        r'<(red|green|blue|orange|underline)>(.*?)</\1>'
        r'|\*\*(.+?)\*\*'
        r'|\*(.+?)\*'
        r'|([^<*]+)'
        r'|(.)',  # fallback cho ký tự lẻ
        re.DOTALL
    )
    
    for m in pattern.finditer(text):
        if m.group(1):  # Color/underline tag
            tag = m.group(1)
            inner = m.group(2)
            if tag == 'underline':
                for sub in _parse_bold_italic(inner):
                    sub['underline'] = True
                    segments.append(sub)
            else:
                color = COLOR_MAP[tag]
                for sub in _parse_bold_italic(inner):
                    sub['color'] = color
                    segments.append(sub)
        elif m.group(3):  # **bold**
            segments.append({
                'text': m.group(3), 'color': None,
                'bold': True, 'italic': False, 'underline': False
            })
        elif m.group(4):  # *italic*
            segments.append({
                'text': m.group(4), 'color': None,
                'bold': False, 'italic': True, 'underline': False
            })
        elif m.group(5):  # plain text
            segments.append({
                'text': m.group(5), 'color': None,
                'bold': False, 'italic': False, 'underline': False
            })
        elif m.group(6):  # fallback single char
            segments.append({
                'text': m.group(6), 'color': None,
                'bold': False, 'italic': False, 'underline': False
            })
    
    return segments


def _parse_bold_italic(text):
    """Parse **bold** and *italic* inside a color/underline tag."""
    segments = []
    pattern = re.compile(r'\*\*(.+?)\*\*|\*(.+?)\*|([^*]+)', re.DOTALL)
    for m in pattern.finditer(text):
        if m.group(1):
            segments.append({
                'text': m.group(1), 'color': None,
                'bold': True, 'italic': False, 'underline': False
            })
        elif m.group(2):
            segments.append({
                'text': m.group(2), 'color': None,
                'bold': False, 'italic': True, 'underline': False
            })
        elif m.group(3):
            segments.append({
                'text': m.group(3), 'color': None,
                'bold': False, 'italic': False, 'underline': False
            })
    return segments


def add_runs_to_paragraph(para, segments, default_color=COLOR_BLACK):
    """Add formatted runs to a paragraph from parsed segments."""
    for seg in segments:
        run = para.add_run(seg['text'])
        color = seg.get('color') or default_color
        run.font.color.rgb = color
        if seg.get('bold'):
            run.bold = True
        if seg.get('italic'):
            run.italic = True
        if seg.get('underline'):
            run.underline = True


def convert_md_to_docx_python_docx(md_path, docx_path):
    """Conversion function using python-docx library."""
    doc = docx.Document()
    
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(14)
    
    for section in doc.sections:
        section.top_margin = Emu(720090)
        section.bottom_margin = Emu(720090)
        section.left_margin = Emu(720090)
        section.right_margin = Emu(720090)
    
    out_dir = os.path.dirname(os.path.abspath(docx_path))
    if out_dir and not os.path.exists(out_dir):
        os.makedirs(out_dir, exist_ok=True)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\n')
        
        if not line.strip():
            i += 1
            continue
        
        if line.strip() == '---':
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(DIVIDER_CHAR)
            run.font.color.rgb = COLOR_DIVIDER
            i += 1
            continue
        
        if line.startswith('## '):
            heading_text = line[3:].strip()
            p = doc.add_paragraph()
            p.add_run(heading_text)
            i += 1
            continue
        
        if line.startswith('# '):
            title_text = line[2:].strip()
            p = doc.add_paragraph()
            p.add_run(title_text)
            i += 1
            continue
        
        if line.startswith('> '):
            bq_text = line[2:].strip()
            p = doc.add_paragraph()
            segments = parse_inline(bq_text)
            add_runs_to_paragraph(p, segments)
            i += 1
            continue
        
        p = doc.add_paragraph()
        segments = parse_inline(line)
        add_runs_to_paragraph(p, segments)
        
        i += 1
    
    doc.save(docx_path)
    print(f'[OK] Saved: {docx_path}')


def convert_md_to_docx_fallback(md_path, docx_path):
    """Pure Python docx converter fallback using zipfile + XML."""
    import zipfile
    import glob

    out_dir = os.path.dirname(os.path.abspath(docx_path))
    if out_dir and not os.path.exists(out_dir):
        os.makedirs(out_dir, exist_ok=True)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    p_xmls = []

    for line in lines:
        line = line.rstrip('\n')
        if not line.strip():
            continue

        if line.strip() == '---':
            p_xml = f'<w:p><w:pPr><w:jc w:val="center"/></w:pPr><w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="28"/><w:szCs w:val="28"/><w:color w:val="C8C8C8"/></w:rPr><w:t xml:space="preserve">{html.escape(DIVIDER_CHAR)}</w:t></w:r></w:p>'
            p_xmls.append(p_xml)
            continue

        if line.startswith('## '):
            line = line[3:].strip()
        elif line.startswith('# '):
            line = line[2:].strip()
        elif line.startswith('> '):
            line = line[2:].strip()

        segments = parse_inline(line)
        runs_xml = []
        for seg in segments:
            text = html.escape(seg['text'])
            color = seg.get('color')
            if color:
                if hasattr(color, 'rgb'):
                    hex_color = f"{color.rgb[0]:02X}{color.rgb[1]:02X}{color.rgb[2]:02X}"
                elif isinstance(color, tuple):
                    hex_color = f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"
                else:
                    hex_color = str(color)
            else:
                hex_color = "000000"

            b_xml = "<w:b/>" if seg.get('bold') else ""
            i_xml = "<w:i/>" if seg.get('italic') else ""
            u_xml = '<w:u w:val="single"/>' if seg.get('underline') else ""

            run_xml = f'<w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="28"/><w:szCs w:val="28"/><w:color w:val="{hex_color}"/>{b_xml}{i_xml}{u_xml}</w:rPr><w:t xml:space="preserve">{text}</w:t></w:r>'
            runs_xml.append(run_xml)

        p_xmls.append(f"<w:p>{''.join(runs_xml)}</w:p>")

    doc_body = "".join(p_xmls)

    document_xml = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
            xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"
            xmlns:v="urn:schemas-microsoft-com:vml"
            xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
            xmlns:w10="urn:schemas-microsoft-com:office:word"
            xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"
            xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml">
  <w:body>
    {doc_body}
    <w:sectPr>
      <w:pgMar w:top="720" w:right="720" w:bottom="720" w:left="720"/>
    </w:sectPr>
  </w:body>
</w:document>'''

    template_docx_path = None
    target_abs = os.path.abspath(docx_path)
    project_root = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
    template_glob = os.path.join(project_root, 'outputs', '**', '*.docx')
    existing_docxs = [d for d in glob.glob(template_glob, recursive=True) if os.path.abspath(d) != target_abs]
    if existing_docxs:
        template_docx_path = existing_docxs[0]

    out_dir = os.path.dirname(docx_path)
    if out_dir and not os.path.exists(out_dir):
        os.makedirs(out_dir, exist_ok=True)

    if template_docx_path and os.path.exists(template_docx_path):
        with zipfile.ZipFile(template_docx_path, 'r') as zin:
            with zipfile.ZipFile(docx_path, 'w', zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    if item.filename == 'word/document.xml':
                        zout.writestr(item.filename, document_xml.encode('utf-8'))
                    else:
                        zout.writestr(item.filename, zin.read(item.filename))
    else:
        content_types_xml = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>'''
        rels_xml = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>'''
        with zipfile.ZipFile(docx_path, 'w', zipfile.ZIP_DEFLATED) as zout:
            zout.writestr('[Content_Types].xml', content_types_xml.encode('utf-8'))
            zout.writestr('_rels/.rels', rels_xml.encode('utf-8'))
            zout.writestr('word/document.xml', document_xml.encode('utf-8'))

    print(f'[OK] Saved (Fallback): {docx_path}')


def convert_md_to_docx(md_path, docx_path):
    """Main conversion entry point."""
    md_path = os.path.abspath(md_path)
    docx_path = os.path.abspath(docx_path)

    if not os.path.exists(md_path):
        raise FileNotFoundError(f'Input markdown not found: {md_path}')

    if HAS_DOCX:
        convert_md_to_docx_python_docx(md_path, docx_path)
    else:
        convert_md_to_docx_fallback(md_path, docx_path)


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print('Usage: python md_to_docx_v2.py <input.md> [output.docx]')
        sys.exit(1)
    
    in_file = sys.argv[1]
    if len(sys.argv) >= 3:
        out_file = sys.argv[2]
    else:
        out_file = os.path.splitext(in_file)[0] + '.docx'
        
    convert_md_to_docx(in_file, out_file)


